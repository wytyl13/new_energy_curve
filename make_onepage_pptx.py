# 一页原生形状版流程图（参考 algorithm_design_base.pptx 模板风格 + 五阶段配色）
# 坐标系 0-100 × 0-100，自上而下，按 16:9 映射到 13.333×7.5 英寸
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

W, H = 13.333, 7.5
INK = RGBColor(0x33, 0x33, 0x33)
ARROW_C = RGBColor(0x6a, 0x6a, 0x6a)
STAGES = {  # 阶段: (填充, 描边)
    '输入':   (RGBColor(0xD6, 0xE4, 0xF7), RGBColor(0x2E, 0x5E, 0x9E)),
    '预处理': (RGBColor(0xD9, 0xF0, 0xE9), RGBColor(0x1F, 0x7A, 0x5C)),
    '九库':   (RGBColor(0xFC, 0xEF, 0xD3), RGBColor(0xB0, 0x7A, 0x1F)),
    '出力':   (RGBColor(0xEB, 0xE2, 0xF6), RGBColor(0x6C, 0x48, 0xA8)),
    '输出':   (RGBColor(0xFA, 0xE3, 0xE5), RGBColor(0xA6, 0x3A, 0x45)),
}
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
S = prs.slides.add_slide(prs.slide_layouts[6])


def set_font(run, size, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, 'Microsoft YaHei'
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', 'Microsoft YaHei')


def text_at(x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.CENTER):
    tb = S.shapes.add_textbox(Inches(x * W / 100), Inches(y * H / 100),
                              Inches(w * W / 100), Inches(h * H / 100))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    set_font(r, size, bold, color)


def box(x, y, w, h, title, lines, stage, dashed=False, tfs=7.0, bfs=5.7):
    fc, ec = STAGES[stage]
    sh = S.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(x * W / 100), Inches(y * H / 100),
                            Inches(w * W / 100), Inches(h * H / 100))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if dashed else fc
    sh.line.color.rgb = ec
    sh.line.width = Pt(1.0)
    if dashed:
        sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    set_font(r, tfs, True, ec)
    for line in lines:
        p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line
        set_font(r, bfs, False, INK)
    return sh


def pill(x, y, w, h, text, stage):
    _, ec = STAGES[stage]
    sh = S.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(x * W / 100), Inches(y * H / 100),
                            Inches(w * W / 100), Inches(h * H / 100))
    sh.fill.solid(); sh.fill.fore_color.rgb = ec
    sh.line.color.rgb = ec; sh.line.width = Pt(0.5)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    set_font(r, 8, True, RGBColor(0xFF, 0xFF, 0xFF))


def conn(x1, y1, x2, y2, head=True):
    c = S.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                               Inches(x1 * W / 100), Inches(y1 * H / 100),
                               Inches(x2 * W / 100), Inches(y2 * H / 100))
    c.line.color.rgb = ARROW_C
    c.line.width = Pt(1.2)
    c.shadow.inherit = False
    if head:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:tailEnd'),
                                 {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c


def elbow(pts):
    for (a, b) in zip(pts[:-2], pts[1:-1]):
        conn(a[0], a[1], b[0], b[1], head=False)
    conn(pts[-2][0], pts[-2][1], pts[-1][0], pts[-1][1])


# ---- 标题 ----
text_at(0, 0.8, 100, 3, '海南风光九条出力曲线流水线：76年NC → 九条8760小时出力标幺曲线', 13, True)
text_at(0, 4.2, 100, 2.4, '每框一个操作，框内标注处理后数据与维度；时间轴删闰年2月29日，每年8760小时', 7.5,
        color=RGBColor(0x66, 0x66, 0x66))

# ---- 阶段标签 ----
pill(1.0, 8.6, 14.5, 3.8, '输入', '输入')
pill(18.5, 8.6, 17.5, 3.8, '预处理与缓存', '预处理')
pill(39.0, 8.6, 27.0, 3.8, '分类 · 代表年 · 插入（9库）', '九库')
pill(68.5, 8.6, 14.5, 3.8, '出力换算', '出力')
pill(85.5, 8.6, 13.5, 3.8, '输出', '输出')

# ---- C1 输入 ----
box(1.0, 16, 14.5, 14, '光NC 76个年文件', ['ssrd (8760,8,10) J/m²·h', '（逐小时累计）'], '输入', tfs=7.3)
box(1.0, 34, 14.5, 14, '风NC 76个年文件', ['u100、v100 (8760,8,10) m/s'], '输入', tfs=7.3)
box(1.0, 52, 14.5, 24, '固定输入（虚线=辅助）', [
    '80格海陆分类表（31海/49陆）', '海南省边界 GeoJSON',
    '00_check_grid_mask：', '边界逐格核验 → 核对.json'], '输入', dashed=True)

# ---- C2 预处理与缓存（01光左上、01风右下错开，箭头互不穿框）----
box(18.5, 16, 14.5, 18, '01.load_one_resource（光）', [
    '80格等权均值÷3600', '→(8760,) W/m²', '删2/29，验证每年8760h'], '预处理')
box(21.0, 36, 15.0, 18, '01.load_one_resource（风）', [
    'hypot(u100,v100)→(8760,80)', '海31格/陆49格等权均值', '→(8760,2)',
    '删2/29，验证每年8760h'], '预处理')
box(18.5, 60, 17.5, 24, '01.main 存NPZ ＋ 02.load_cache', [
    '风速152键 {年}_sea/{年}_land', '光伏76键 {年}，各(8760,)',
    '→ 海风/陆风/光各 (76,8760)', '光负噪声≥−0.01截为0'], '预处理')
conn(15.5, 23, 18.5, 25)
conn(15.5, 41, 21.0, 45)
conn(20.0, 34, 20.0, 60)
conn(34.0, 54, 34.0, 60)

# ---- C3 分类·代表年·插入 ----
box(39.0, 16, 27.0, 12, '02.classify_years', [
    '年指标：风 mean(v³)、光 ΣG/1000', '76年升序分组：小19 / 平均38 / 大19'], '九库')
PAR = [
    ('海风库 ×3', ['X (N,8760)，N=19/38', 'medoid → 真实代表年', 'before (8760,)',
                'insert_wind', 'q10/q90 → 候选池', '逐月批次插入', 'after (8760,)']),
    ('陆风库 ×3', ['X (N,8760)，N=19/38', 'medoid → 真实代表年', 'before (8760,)',
                'insert_wind', 'q10/q90 → 候选池', '逐月批次插入', 'after (8760,)']),
    ('光伏库 ×3', ['X (N,8760)，N=19/38', 'medoid → 真实代表年', 'before (8760,)',
                'insert_solar', '(N,365,24)→日均(N,365)', '低光模板(24,)整日替换', 'after (8760,)']),
]
PXS = [39.0, 48.4, 57.8]
for (t, ls), px in zip(PAR, PXS):
    box(px, 34, 8.2, 28, t, ls, '九库', tfs=6.3, bfs=5.2)
    conn(px + 4.1, 28, px + 4.1, 34)
    conn(px + 4.1, 62, px + 4.1, 68)
box(39.0, 68, 27.0, 12, '02.export_group 写库', [
    '9 × curve_8760.csv (8760,)', '各配 summary / candidates / insertion_log'], '九库')

# ---- C4 出力换算 ----
box(68.5, 30, 14.5, 28, '03_power_curves.main', [
    '（9支路合并）', 'wind_power：', '海风 4/11/25 α=0.70', '陆风 2.5/9/18 α=0.95',
    '光伏 p=0.6·G/1000', '资源(8760,)→pu(8760,)', '插前/插后 ×9'], '出力')

# ---- C5 输出 ----
box(85.5, 18, 13.5, 14, '九条出力总表', ['(8760, 2+9×4)'], '输出')
box(85.5, 36, 13.5, 12, '最大利用小时', ['H=Σp (9,)'], '输出')
box(85.5, 52, 13.5, 13, '04_verify_outputs', ['复核结果.json', '（原极端保护等）'], '输出')
box(85.5, 69, 13.5, 13, '05_plot_results', ['九条曲线 PNG', '（3×3子图）'], '输出')
conn(92.25, 32, 92.25, 36)
conn(92.25, 48, 92.25, 52)
conn(92.25, 65, 92.25, 69)

# ---- 跨列主箭头 ----
elbow([(36.0, 72), (37.6, 72), (37.6, 22), (39.0, 22)])   # NPZ → classify
elbow([(66.0, 74), (67.3, 74), (67.3, 44), (68.5, 44)])   # 写库 → 03
conn(83.0, 36, 85.5, 25)                                   # 03 → 总表
conn(83.0, 44, 85.5, 42)                                   # 03 → H

# ---- 脚注 ----
text_at(0, 94, 100, 3, '风插入：先P10后P90逐月批次，完整片段优先、两侧隔3h、同日至多1条；'
        '光插入：月缺口最大优先，低光日整日24h替换 —— 循环细节见 DESIGN.md',
        6.5, color=RGBColor(0x88, 0x88, 0x88))

prs.save('算法处理流程_单页.pptx')
print('saved 算法处理流程_单页.pptx（原生形状一页版）')
