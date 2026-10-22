# 生成 算法处理流程.pptx：仿 algorithm_design_base.pptx 的框图格式（框=操作，箭头间标数据维度）
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

INK = RGBColor(0x21, 0x21, 0x21)   # 正文
GRAY = RGBColor(0x8a, 0x8a, 0x8a)  # 维度标注
LINE = RGBColor(0x9a, 0x9a, 0x9a)  # 框线
FILL = RGBColor(0xff, 0xff, 0xff)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_font(run, size, bold=False, color=INK, name='Microsoft YaHei'):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, name
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def add_slide_title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.15), Inches(12.4), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    set_font(r, 16, bold=True)


def box(slide, x, y, w, h, title, lines=(), fill=FILL):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE; sh.line.width = Pt(1.0)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    set_font(r, 12, bold=True)
    for line in lines:
        p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line
        set_font(r, 10.5)
    return sh


def arrow(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor(0x55, 0x55, 0x55)
    c.line.width = Pt(1.5)
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return c


def dim_label(slide, x, y, w, text, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.25))
    p = tb.text_frame.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    set_font(r, 10, color=GRAY)
    return tb


# ---------------- Slide 1：输入层 → 76年NPZ缓存 ----------------
s = prs.slides.add_slide(BLANK)
add_slide_title(s, '海南风光九条出力曲线流水线（1/2）：NC → 76年缓存')
CX = 6.667
box(s, 3.67, 0.75, 6.0, 1.05, '输入层', [
    '风NC 76个年文件：u100、v100 (8760,8,10) m/s',
    '光NC 76个年文件：ssrd (8760,8,10) J/m²·h',
    '80格海陆分类表（31海 / 49陆）'])
arrow(s, CX, 1.80, CX, 2.15)
box(s, 3.67, 2.15, 6.0, 0.62, '00_check_grid_mask（模块级）',
    ['边界多边形逐格包含判定 ↔ 分类表 → 网格海陆核对.json（80格判定）'])
arrow(s, CX, 2.77, CX, 3.12)
dim_label(s, 6.85, 2.80, 6.0, '(8760,8,10) × 76年，逐块读入')
# 风光两条并行支路
arrow(s, 3.9, 3.12, 3.9, 3.55)
arrow(s, 9.4, 3.12, 9.4, 3.55)
dim_label(s, 0.15, 3.22, 3.4, '风：每年 (8760,2)', PP_ALIGN.RIGHT)
dim_label(s, 9.6, 3.22, 3.4, '光：每年 (8760,)')
box(s, 0.5, 3.55, 5.6, 1.15, '01_nc_to_cache.load_one_resource（风）', [
    'hypot(u100,v100) → 逐时风速 (8760,80)',
    '海31格等权均值 / 陆49格等权均值 → (8760,2)',
    '删2月29日，验证每年完整无重复'])
box(s, 7.2, 3.55, 5.6, 1.15, '01_nc_to_cache.load_one_resource（光）', [
    '80格等权均值 ÷ 3600 → (8760,) W/m²',
    '删2月29日，验证每年完整无重复'])
arrow(s, 3.9, 4.70, CX, 5.10)
arrow(s, 9.4, 4.70, CX, 5.10)
box(s, 3.67, 5.10, 6.0, 1.05, '01_nc_to_cache.main（存NPZ缓存）', [
    '风速NPZ 152键 {年}_sea/{年}_land 各(8760,)',
    '光伏NPZ 76键 {年} 各(8760,)',
    '→ 海风(76,8760)、陆风(76,8760)、光(76,8760)'])

# ---------------- Slide 2：9库 → 出力 → 复核 ----------------
s = prs.slides.add_slide(BLANK)
add_slide_title(s, '海南风光九条出力曲线流水线（2/2）：分类 → 代表年 → 插入 → 出力')
box(s, 3.67, 0.60, 6.0, 0.66, '02_classify_cluster_insert.load_cache', [
    'NPZ → dict{年:(8760,)} → (76,8760)；光负噪声≥−0.01截为0'])
arrow(s, CX, 1.26, CX, 1.58)
dim_label(s, 6.85, 1.26, 6.0, '(76,8760) × 3条支路')
box(s, 3.67, 1.58, 6.0, 0.70, '02.classify_years', [
    '年指标：风 mean(v³)、光 ΣG/1000 → 升序分组：小19 / 平均38 / 大19'])
COLY, COLH, COLW = 2.85, 1.85, 3.7
XS = [0.9, 4.82, 8.74]
for x in XS:
    arrow(s, CX, 2.28, x + COLW / 2, COLY)
box(s, XS[0], COLY, COLW, COLH, '海风库 ×3（小/平均/大）', [
    'export_group：X (N,8760)，N=19或38',
    'medoid → 真实代表年 before (8760,)',
    'insert_wind：q10/q90 → 候选池 →',
    '逐月批次插入 → after (8760,)'])
box(s, XS[1], COLY, COLW, COLH, '陆风库 ×3（小/平均/大）', [
    'export_group：X (N,8760)，N=19或38',
    'medoid → 真实代表年 before (8760,)',
    'insert_wind：q10/q90 → 候选池 →',
    '逐月批次插入 → after (8760,)'])
box(s, XS[2], COLY, COLW, COLH, '光伏库 ×3（小/平均/大）', [
    'export_group：X (N,8760)，N=19或38',
    'medoid → 真实代表年 before (8760,)',
    'insert_solar：(N,365,24)→日均(N,365)',
    '→低光模板(24,)→整日替换 after (8760,)'])
for x in XS:
    arrow(s, x + COLW / 2, COLY + COLH, CX, 5.05)
box(s, 3.67, 5.05, 6.0, 0.62, '02.export_group（写库）', [
    '9 × curve_8760.csv (8760,)，各配 summary / candidates / insertion_log'])
arrow(s, CX, 5.67, CX, 5.98)
dim_label(s, 6.85, 5.68, 6.0, '(8760,) × 9')
box(s, 3.17, 5.98, 7.0, 0.95, '03_power_curves.main（9条支路在此合并）', [
    'wind_power：海风4/11/25 m/s α=0.70、陆风2.5/9/18 α=0.95；光伏 p=0.6·G/1000',
    '资源 (8760,) → 出力pu (8760,) 插前/插后 ×9；总表 (8760, 2+9×4)；H=Σpₜ (9,)'])
arrow(s, CX, 6.93, CX, 7.14)
box(s, 3.17, 7.14, 7.0, 0.32, '04_verify_outputs：复核结果.json　→　05_plot_results：PNG（3×3子图）', ())

# ---------------- Slide 3：两个插入循环 ----------------
s = prs.slides.add_slide(BLANK)
add_slide_title(s, '步骤2 内部：风插入与光插入的循环结构')
wind_loop = '\n'.join([
    'insert_wind（风库，先P10后P90，共用同一曲线）',
    '',
    '循环外（一次）：算P10/P90 → 候选池（qualified_wind，',
    '  越阈且连续≥3h片段）→ 全年/月目标（库内逐年越阈小时',
    '  均值，half_up）→ 记录原越阈小时',
    'for 尾部 in (P10, P90)：',
    '  while 全年剩余≥3 且 有活跃月：',
    '    重算越阈小时 → 全年剩余、12个月缺口',
    '    选缺口最大月（同缺口先早月）',
    '    额度 = 缺口>6 ? ceil(缺口/2) : 3，且≤全年剩余',
    '    try_wind_batch：完整片段优先；均不可行才允许前缀',
    '      截断（截后≥3h）；两侧各隔3个正常小时；同日至多',
    '      1条；达不到额度整批不提交',
    '    成立→提交（更新曲线/占用日/日志）；否则冻结该月',
    '  自检：原越阈小时不变、片段不覆盖、目标位之外不变'])
solar_loop = '\n'.join([
    'insert_solar（光库）',
    '',
    '循环外（一次）：低光判定（日均≤P10）→ 月目标天=库内',
    '  年均值；全年目标天=库内年均值 且 ≥ 代表年原有天数 →',
    '  每个可插日的模板（同月同日历史低光日按日总量排序，',
    '  奇数取中间1天、偶数取中间2天逐时平均；剔除代表年',
    '  原有低光日）',
    'while 全年剩余>0 且 有活跃月：',
    '  重算低光日 → 月缺口，选缺口最大月，额度=ceil(缺口/2)',
    '  该月按历史低光频次选日，整日24h替换为模板',
    '  额度不足或无候选 → 冻结该月',
    '自检：原有低光日、未选日期均未被改动'])
for x, text in ((0.5, wind_loop), (6.85, solar_loop)):
    tb = s.shapes.add_textbox(Inches(x), Inches(0.7), Inches(6.0), Inches(6.5))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line if line else ' '
        set_font(r, 11, bold=(i == 0))

prs.save('算法处理流程.pptx')
print('saved 算法处理流程.pptx')
