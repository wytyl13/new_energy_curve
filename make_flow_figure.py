# 生成 算法流程图.png（单张全流程图，按阶段配色）+ 单页pptx
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch

plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False

# 阶段配色：填充 / 描边
STAGES = {
    '输入': ('#D6E4F7', '#2E5E9E'),
    '预处理与缓存': ('#D9F0E9', '#1F7A5C'),
    '分类·代表年·插入': ('#FCEFD3', '#B07A1F'),
    '出力换算': ('#EBE2F6', '#6C48A8'),
    '输出': ('#FAE3E5', '#A63A45'),
}
ARROW = '#6a6a6a'

fig, ax = plt.subplots(figsize=(16, 9), dpi=200)
ax.set_xlim(0, 100); ax.set_ylim(0, 100); ax.axis('off')
ax.set_facecolor('white')


def box(x, y, w, h, title, lines, stage, dashed=False, tfs=8.2, bfs=6.9):
    fc, ec = STAGES[stage]
    p = FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0,rounding_size=1.1',
                       fc=('white' if dashed else fc), ec=ec,
                       lw=1.2, ls=('--' if dashed else '-'))
    ax.add_patch(p)
    ys = np.linspace(y + h - 2.2, y + 1.0, len(lines) + 1)
    ax.text(x + w / 2, ys[0], title, ha='center', va='center',
            fontsize=tfs, fontweight='bold', color=ec)
    for yi, line in zip(ys[1:], lines):
        ax.text(x + w / 2, yi, line, ha='center', va='center',
                fontsize=bfs, color='#333333')


def pill(x, y, w, h, text, stage):
    _, ec = STAGES[stage]
    p = FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0,rounding_size=0.9',
                       fc=ec, ec=ec, lw=0)
    ax.add_patch(p)
    ax.text(x + w / 2, y + h / 2, text, ha='center', va='center',
            fontsize=9.5, fontweight='bold', color='white')


def arrow(p1, p2, color=ARROW, lw=1.3):
    ax.annotate('', xy=p2, xytext=p1,
                arrowprops=dict(arrowstyle='-|>', color=color, lw=lw,
                                mutation_scale=13, shrinkA=0, shrinkB=0))


def elbow(points, color=ARROW, lw=1.3):
    for a, b in zip(points[:-2], points[1:-1]):
        ax.plot([a[0], b[0]], [a[1], b[1]], color=color, lw=lw,
                solid_capstyle='round')
    arrow(points[-2], points[-1], color, lw)


# ---- 标题 ----
ax.text(50, 97.5, '海南风光九条出力曲线流水线：76年NC → 九条8760小时出力标幺曲线',
        ha='center', va='center', fontsize=16, fontweight='bold')
ax.text(50, 93.3, '每框一个操作，框内标注处理后数据与维度；时间轴删闰年2月29日，每年8760小时',
        ha='center', va='center', fontsize=9, color='#666666')

# ---- 阶段标签 ----
pill(1.0, 89.0, 14.5, 3.6, '输入', '输入')
pill(18.5, 89.0, 17.5, 3.6, '预处理与缓存', '预处理与缓存')
pill(39.0, 89.0, 27.0, 3.6, '分类 · 代表年 · 插入（9库）', '分类·代表年·插入')
pill(68.5, 89.0, 14.5, 3.6, '出力换算', '出力换算')
pill(85.5, 89.0, 13.5, 3.6, '输出', '输出')

# ---- C1 输入 ----
box(1.0, 72, 14.5, 14, '光NC 76个年文件', ['ssrd (8760,8,10) J/m²·h', '（逐小时累计）'], '输入', tfs=8.6)
box(1.0, 50, 14.5, 14, '风NC 76个年文件', ['u100、v100 (8760,8,10) m/s'], '输入', tfs=8.6)
box(1.0, 20, 14.5, 24, '固定输入（虚线=辅助）', [
    '80格海陆分类表（31海/49陆）', '海南省边界 GeoJSON',
    '00_check_grid_mask：', '边界逐格核验 → 核对.json'], '输入', dashed=True)

# ---- C2 预处理与缓存（01光左上、01风右下错开，箭头互不穿框）----
box(18.5, 66, 14.5, 18, '01.load_one_resource（光）', [
    '80格等权均值÷3600', '→(8760,) W/m²', '删2/29，验证每年8760h'], '预处理与缓存')
box(21.0, 44, 15.0, 18, '01.load_one_resource（风）', [
    'hypot(u100,v100)→(8760,80)', '海31格/陆49格等权均值',
    '→(8760,2)', '删2/29，验证每年8760h'], '预处理与缓存')
box(18.5, 12, 17.5, 24, '01.main 存NPZ ＋ 02.load_cache', [
    '风速152键 {年}_sea/{年}_land', '光伏76键 {年}，各(8760,)',
    '→ 海风/陆风/光各 (76,8760)', '光负噪声≥−0.01截为0'], '预处理与缓存')
arrow((15.5, 79), (18.5, 75))
arrow((15.5, 57), (21.0, 53))
arrow((20.0, 66), (20.0, 36))
arrow((34.0, 44), (34.0, 36))

# ---- C3 分类·代表年·插入 ----
box(39.0, 70, 27.0, 12, '02.classify_years', [
    '年指标：风 mean(v³)、光 ΣG/1000', '76年升序分组：小19 / 平均38 / 大19'], '分类·代表年·插入')
PAR = [
    ('海风库 ×3', ['X (N,8760)，N=19/38', 'medoid → 真实代表年', 'before (8760,)',
                'insert_wind', 'q10/q90 → 候选池', '逐月批次插入', 'after (8760,)']),
    ('陆风库 ×3', ['X (N,8760)，N=19/38', 'medoid → 真实代表年', 'before (8760,)',
                'insert_wind', 'q10/q90 → 候选池', '逐月批次插入', 'after (8760,)']),
    ('光伏库 ×3', ['X (N,8760)，N=19/38', 'medoid → 真实代表年', 'before (8760,)',
                'insert_solar', '(N,365,24)→日均(N,365)', '低光模板(24,)整日替换', 'after (8760,)']),
]
PXS = [39.0, 48.4, 57.8]
for (t, ls), px in zip(PAR, PXS):
    box(px, 36, 8.2, 28, t, ls, '分类·代表年·插入', tfs=7.6, bfs=6.2)
    arrow((px + 4.1, 70), (px + 4.1, 64))
    arrow((px + 4.1, 36), (px + 4.1, 24))
box(39.0, 12, 27.0, 12, '02.export_group 写库', [
    '9 × curve_8760.csv (8760,)', '各配 summary / candidates / insertion_log'], '分类·代表年·插入')

# ---- C4 出力换算 ----
box(68.5, 44, 14.5, 28, '03_power_curves.main', [
    '（9支路合并）', 'wind_power：', '海风 4/11/25 α=0.70', '陆风 2.5/9/18 α=0.95',
    '光伏 p=0.6·G/1000', '资源(8760,)→pu(8760,)', '插前/插后 ×9'], '出力换算')

# ---- C5 输出 ----
box(85.5, 63, 13.5, 14, '九条出力总表', ['(8760, 2+9×4)'], '输出')
box(85.5, 47, 13.5, 12, '最大利用小时', ['H=Σp (9,)'], '输出')
box(85.5, 30, 13.5, 13, '04_verify_outputs', ['复核结果.json', '（原极端保护等）'], '输出')
box(85.5, 13, 13.5, 13, '05_plot_results', ['九条曲线 PNG', '（3×3子图）'], '输出')
arrow((85.5, 58), (85.5 + 6.75, 47 + 6))
arrow((85.5 + 6.75, 47), (85.5 + 6.75, 43))
arrow((85.5 + 6.75, 30), (85.5 + 6.75, 26))

# ---- 跨列主箭头 ----
elbow([(36.0, 24), (37.6, 24), (37.6, 76), (39.0, 76)])          # NPZ → classify
elbow([(66.0, 18), (67.3, 18), (67.3, 58), (68.5, 58)])          # 写库 → 03
arrow((83.0, 62), (85.5, 70))                                     # 03 → 总表
arrow((83.0, 54), (85.5, 53))                                     # 03 → H

# ---- 脚注 ----
ax.text(50, 6.5, '风插入：先P10后P90逐月批次，完整片段优先、两侧隔3h、同日至多1条；'
        '光插入：月缺口最大优先，低光日整日24h替换 —— 循环细节见 DESIGN.md',
        ha='center', va='center', fontsize=8, color='#888888')

fig.savefig('算法流程图.png', bbox_inches='tight', facecolor='white')
print('saved 算法流程图.png')

# ---- 单页pptx ----
from pptx import Presentation
from pptx.util import Inches
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.shapes.add_picture('算法流程图.png', 0, 0, width=prs.slide_width, height=prs.slide_height)
prs.save('算法处理流程_单页.pptx')
print('saved 算法处理流程_单页.pptx')
